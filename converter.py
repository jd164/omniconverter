"""
Multi-Format Call Log Conversion Engine.
Supported formats: Excel (.xlsx), CSV (.csv), JSON (.json), and iCalendar (.ics).
"""

from datetime import datetime, timedelta
import json
import os
import re
import uuid
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def format_duration(seconds: int) -> str:
    """Formats seconds into HH:MM:SS or MM:SS."""
    if seconds is None:
        return "00:00"
    seconds = int(seconds)
    hours = seconds // 3600
    minutes = (seconds % 3600) // 60
    secs = seconds % 60
    if hours > 0:
        return f"{hours:02d}:{minutes:02d}:{secs:02d}"
    return f"{minutes:02d}:{secs:02d}"


def load_data(source, limit=None, selected_ids=None):
    """
    Loads and normalizes call log data from file path or dictionary.
    Optionally filters by limit or specific selected IDs.
    """
    if isinstance(source, str):
        if not os.path.exists(source):
            for alt in ["data.json", "dados.json", "sample_data.json"]:
                if os.path.exists(alt):
                    source = alt
                    break
            if not os.path.exists(source):
                raise FileNotFoundError(f"File not found: {source}")
        with open(source, "r", encoding="utf-8") as f:
            raw = json.load(f)
    elif isinstance(source, dict):
        raw = source
    elif isinstance(source, list):
        raw = {"calls": source}
    else:
        raise ValueError("Unsupported data format. Expected file path, dict, or list.")

    calls_raw = raw.get("calls", [])
    contact = raw.get("contact", {})
    exported_at = raw.get("exportedAt", datetime.utcnow().isoformat() + "Z")

    normalized_calls = []
    for c in calls_raw:
        call_id = c.get("id", str(uuid.uuid4()))
        if selected_ids and call_id not in selected_ids:
            continue

        duration_sec = c.get("durationSec", 0)
        dur_formatted = c.get("durationFormatted") or format_duration(duration_sec)

        iso_str = c.get("iso", "")
        date_str = c.get("date", "")
        time_str = c.get("time", "")

        if not iso_str and date_str and time_str:
            iso_str = f"{date_str}T{time_str}.000Z"
        elif iso_str and (not date_str or not time_str):
            try:
                dt = datetime.fromisoformat(iso_str.replace("Z", "+00:00"))
                date_str = dt.strftime("%Y-%m-%d")
                time_str = dt.strftime("%H:%M:%S")
            except Exception:
                pass

        media_type = c.get("mediaType", "Voice")
        if media_type.lower() in ("voz", "voice", "audio"):
            media_type = "Voice"
        elif media_type.lower() in ("vídeo", "video"):
            media_type = "Video"

        from_me = c.get("fromMe", False)
        direction = "Outgoing" if from_me else "Incoming"
        status = c.get("status", "Incoming (Answered)" if not from_me else "Outgoing (Answered)")
        is_missed = c.get("isMissed", False)
        status_lower = status.lower()
        if "missed" in status_lower or "perdida" in status_lower or "não atendida" in status_lower or "unanswered" in status_lower:
            is_missed = True

        normalized_calls.append({
            "id": call_id,
            "chatId": c.get("chatId", ""),
            "timestamp": c.get("timestamp", 0),
            "timestampMs": c.get("timestampMs", c.get("timestamp", 0) * 1000),
            "date": date_str,
            "time": time_str,
            "iso": iso_str,
            "mediaType": media_type,
            "direction": direction,
            "status": status,
            "isMissed": is_missed,
            "durationSec": duration_sec,
            "durationFormatted": dur_formatted,
            "fromMe": from_me,
            "source": c.get("source", "whatsapp-call-export")
        })

    if limit and limit > 0:
        normalized_calls = normalized_calls[:limit]

    total_calls = len(normalized_calls)
    total_duration = sum(c["durationSec"] for c in normalized_calls)
    voice_count = sum(1 for c in normalized_calls if c["mediaType"].lower() == "voice")
    video_count = sum(1 for c in normalized_calls if c["mediaType"].lower() == "video")
    missed_count = sum(1 for c in normalized_calls if c["isMissed"])
    answered_count = total_calls - missed_count

    return {
        "exportedAt": exported_at,
        "contact": contact,
        "totalCalls": total_calls,
        "totalDurationSeconds": total_duration,
        "totalDurationFormatted": format_duration(total_duration),
        "averageDurationSeconds": int(total_duration / total_calls) if total_calls > 0 else 0,
        "averageDurationFormatted": format_duration(int(total_duration / total_calls)) if total_calls > 0 else "00:00",
        "voiceCalls": voice_count,
        "videoCalls": video_count,
        "missedCalls": missed_count,
        "answeredCalls": answered_count,
        "calls": normalized_calls
    }


def export_excel(source, output_path="call_logs.xlsx", event_title="Call", limit=None, selected_ids=None):
    data = load_data(source, limit=limit, selected_ids=selected_ids)
    calls = data["calls"]

    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "Call Log"
    ws.views.sheetView[0].showGridLines = True

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10, color="0F172A")
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    thin_border_side = Side(border_style="thin", color="CBD5E1")
    border_cell = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    headers = [
        ("No.", Alignment(horizontal="center")),
        ("Date", Alignment(horizontal="center")),
        ("Time", Alignment(horizontal="center")),
        ("Type", Alignment(horizontal="center")),
        ("Direction", Alignment(horizontal="center")),
        ("Status", Alignment(horizontal="left")),
        ("Duration (s)", Alignment(horizontal="right")),
        ("Formatted Duration", Alignment(horizontal="center")),
        ("UTC Start (ISO)", Alignment(horizontal="center")),
        ("Record ID", Alignment(horizontal="left")),
    ]

    for col_idx, (col_name, align) in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border_cell
    ws.row_dimensions[1].height = 28

    for row_idx, call in enumerate(calls, start=2):
        row_fill = zebra_fill if row_idx % 2 == 0 else white_fill
        valores = [
            (row_idx - 1, Alignment(horizontal="center")),
            (call["date"], Alignment(horizontal="center")),
            (call["time"], Alignment(horizontal="center")),
            (call["mediaType"], Alignment(horizontal="center")),
            (call["direction"], Alignment(horizontal="center")),
            (call["status"], Alignment(horizontal="left")),
            (call["durationSec"], Alignment(horizontal="right")),
            (call["durationFormatted"], Alignment(horizontal="center")),
            (call["iso"], Alignment(horizontal="center")),
            (call["id"], Alignment(horizontal="left")),
        ]

        for col_idx, (val, align) in enumerate(valores, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = data_font
            cell.fill = row_fill
            cell.alignment = align
            cell.border = border_cell
        ws.row_dimensions[row_idx].height = 20

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    ws_stats = wb.create_sheet(title="Statistical Summary")
    ws_stats.views.sheetView[0].showGridLines = True

    title_fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
    title_font = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
    section_font = Font(name="Segoe UI", size=11, bold=True, color="1E293B")
    metric_label_font = Font(name="Segoe UI", size=10, bold=True, color="475569")
    metric_val_font = Font(name="Segoe UI", size=11, bold=True, color="0F172A")

    ws_stats.merge_cells("A1:D1")
    title_cell = ws_stats["A1"]
    title_cell.value = "Call Log Analytics Report"
    title_cell.fill = title_fill
    title_cell.font = title_font
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_stats.row_dimensions[1].height = 36

    contact_name = data.get("contact", {}).get("name", "N/A")
    contact_phone = data.get("contact", {}).get("phone", "N/A")
    min_date = min((c["date"] for c in calls if c["date"]), default="N/A")
    max_date = max((c["date"] for c in calls if c["date"]), default="N/A")

    summary_items = [
        ("General Information", ""),
        ("Contact / Name", contact_name),
        ("Phone Number", contact_phone if contact_phone else "N/A"),
        ("Date Range", f"{min_date} to {max_date}"),
        ("Export Date", data.get("exportedAt", "N/A")),
        ("", ""),
        ("Volume Metrics", ""),
        ("Total Calls", data["totalCalls"]),
        ("Voice Calls", data["voiceCalls"]),
        ("Video Calls", data["videoCalls"]),
        ("Answered Calls", data["answeredCalls"]),
        ("Missed Calls", data["missedCalls"]),
        ("", ""),
        ("Duration Metrics", ""),
        ("Total Duration (seconds)", data["totalDurationSeconds"]),
        ("Total Duration Formatted", data["totalDurationFormatted"]),
        ("Average Duration per Call", data["averageDurationFormatted"]),
    ]

    for idx, (label, val) in enumerate(summary_items, start=3):
        ws_stats.row_dimensions[idx].height = 22
        if label and val == "":
            c1 = ws_stats.cell(row=idx, column=1, value=label)
            c1.font = section_font
            ws_stats.cell(row=idx, column=2, value="")
        elif label == "" and val == "":
            pass
        else:
            c1 = ws_stats.cell(row=idx, column=1, value=label)
            c2 = ws_stats.cell(row=idx, column=2, value=val)
            c1.font = metric_label_font
            c2.font = metric_val_font
            c1.border = border_cell
            c2.border = border_cell
            c1.fill = zebra_fill
            c2.alignment = Alignment(horizontal="right" if isinstance(val, (int, float)) else "left")

    ws_stats.column_dimensions["A"].width = 30
    ws_stats.column_dimensions["B"].width = 35

    wb.save(output_path)
    return output_path


def export_csv(source, output_path="call_logs.csv", delimiter=";", limit=None, selected_ids=None):
    import csv

    data = load_data(source, limit=limit, selected_ids=selected_ids)
    calls = data["calls"]

    with open(output_path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f, delimiter=delimiter, quoting=csv.QUOTE_MINIMAL)
        writer.writerow([
            "ID",
            "Date",
            "Time",
            "ISO_UTC",
            "Media_Type",
            "Direction",
            "Status",
            "Missed",
            "Duration_Seconds",
            "Duration_Formatted",
            "Chat_ID",
            "From_Me"
        ])

        for call in calls:
            writer.writerow([
                call["id"],
                call["date"],
                call["time"],
                call["iso"],
                call["mediaType"],
                call["direction"],
                call["status"],
                "Yes" if call["isMissed"] else "No",
                call["durationSec"],
                call["durationFormatted"],
                call["chatId"],
                "Yes" if call["fromMe"] else "No"
            ])

    return output_path


def export_json(source, output_path="call_logs.json", indent=2, limit=None, selected_ids=None):
    data = load_data(source, limit=limit, selected_ids=selected_ids)
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=indent)
    return output_path


def export_ics(source, output_path="call_logs.ics", event_title="Call with {contact}", limit=None, selected_ids=None):
    data = load_data(source, limit=limit, selected_ids=selected_ids)
    calls = data["calls"]

    now_str = datetime.utcnow().strftime("%Y%m%dT%H%M%SZ")
    clean_cal_title = event_title.title()

    ics_lines = [
        "BEGIN:VCALENDAR",
        "VERSION:2.0",
        "PRODID:-//OmniConverter//EN",
        "CALSCALE:GREGORIAN",
        "METHOD:PUBLISH",
        f"X-WR-CALNAME:{clean_cal_title}",
    ]

    for call in calls:
        iso_str = call.get("iso", "")
        if not iso_str:
            continue

        try:
            clean_iso = iso_str.replace("Z", "").split(".")[0]
            dt_start = datetime.strptime(clean_iso, "%Y-%m-%dT%H:%M:%S")
        except Exception:
            continue

        duration_sec = call.get("durationSec", 0)
        dur_calendar_sec = duration_sec if duration_sec > 0 else 300
        dt_end = dt_start + timedelta(seconds=dur_calendar_sec)

        uid = call.get("id", str(uuid.uuid4()))
        dtstart_str = dt_start.strftime("%Y%m%dT%H%M%SZ")
        dtend_str = dt_end.strftime("%Y%m%dT%H%M%SZ")

        summary = event_title
        contact_name = data.get("contact", {}).get("name", "Contact")
        summary = summary.replace("{contact}", contact_name)
        summary = summary.replace("{type}", call.get("mediaType", "Voice"))
        summary = summary.replace("{duration}", call.get("durationFormatted", ""))

        description = (
            f"Type: {call.get('mediaType', 'Voice')} Call\\n"
            f"Duration: {call.get('durationFormatted', '00:00')}\\n"
            f"Status: {call.get('status', 'Completed')}\\n"
            f"Direction: {call.get('direction', 'Incoming')}"
        )

        ics_lines.extend([
            "BEGIN:VEVENT",
            f"UID:{uid}",
            f"DTSTAMP:{now_str}",
            f"DTSTART:{dtstart_str}",
            f"DTEND:{dtend_str}",
            f"SUMMARY:{summary}",
            f"DESCRIPTION:{description}",
            "STATUS:CONFIRMED",
            "TRANSP:OPAQUE",
            "END:VEVENT",
        ])

    ics_lines.append("END:VCALENDAR")

    with open(output_path, "w", encoding="utf-8") as f:
        f.write("\r\n".join(ics_lines))

    return output_path


def convert_all(source, base_name="call_logs", event_title="Call with {contact}", limit=None, selected_ids=None):
    data = load_data(source, limit=limit, selected_ids=selected_ids)

    xlsx_file = f"{base_name}.xlsx"
    csv_file = f"{base_name}.csv"
    json_file = f"{base_name}.json"
    ics_file = f"{base_name}.ics"

    export_excel(data, xlsx_file, event_title=event_title)
    export_csv(data, csv_file)
    export_json(data, json_file)
    export_ics(data, ics_file, event_title=event_title)

    return {
        "xlsx": xlsx_file,
        "csv": csv_file,
        "json": json_file,
        "ics": ics_file,
        "totalCalls": data["totalCalls"],
        "totalDuration": data["totalDurationFormatted"]
    }


# Backwards compatibility aliases
carregar_dados = load_data
exportar_excel = export_excel
exportar_csv = export_csv
exportar_json = export_json
exportar_ics = export_ics
converter_todos = convert_all
